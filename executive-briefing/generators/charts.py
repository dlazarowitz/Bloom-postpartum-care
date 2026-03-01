"""
Shared chart-building utilities for PowerPoint generation.
Uses python-pptx's native chart API for editable charts.
"""

from pptx.util import Inches, Pt, Emu
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.dml.color import RGBColor

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from templates.styles import hex_to_rgb, get_chart_colors


def add_bar_chart(slide, categories, series_data, config,
                  left=None, top=None, width=None, height=None,
                  chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED,
                  show_legend=True, title=None):
    """
    Add a bar/column chart to a slide.

    Args:
        categories: list of category labels
        series_data: dict of {series_name: [values]}
        config: branding config dict
    """
    from templates.styles import CHART_LEFT, CHART_TOP, CHART_WIDTH, CHART_HEIGHT

    left = left or CHART_LEFT
    top = top or CHART_TOP
    width = width or CHART_WIDTH
    height = height or CHART_HEIGHT

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series_data.items():
        chart_data.add_series(name, values)

    chart_frame = slide.shapes.add_chart(
        chart_type, left, top, width, height, chart_data
    )
    chart = chart_frame.chart
    chart_colors = get_chart_colors(config)

    if title:
        chart.has_title = True
        chart.chart_title.text_frame.paragraphs[0].text = title
        chart.chart_title.text_frame.paragraphs[0].font.size = Pt(12)

    if show_legend:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9)
    else:
        chart.has_legend = False

    # Apply colors to series
    plot = chart.plots[0]
    for i, series in enumerate(plot.series):
        if i < len(chart_colors):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = chart_colors[i]

    # Format value axis
    value_axis = chart.value_axis
    value_axis.has_minor_gridlines = False
    value_axis.major_gridlines.format.line.color.rgb = hex_to_rgb("DDDDDD")
    value_axis.tick_labels.font.size = Pt(9)

    # Format category axis
    category_axis = chart.category_axis
    category_axis.tick_labels.font.size = Pt(9)

    return chart_frame


def add_line_chart(slide, categories, series_data, config,
                   left=None, top=None, width=None, height=None,
                   show_legend=True, title=None):
    """Add a line chart to a slide."""
    return add_bar_chart(
        slide, categories, series_data, config,
        left=left, top=top, width=width, height=height,
        chart_type=XL_CHART_TYPE.LINE_MARKERS,
        show_legend=show_legend, title=title,
    )


def add_table(slide, rows, col_widths, config,
              left=None, top=None, width=None, height=None,
              header_color=None):
    """
    Add a formatted table to a slide.

    Args:
        rows: list of lists [[header1, header2...], [row1col1, row1col2...], ...]
        col_widths: list of Inches values for each column
        config: branding config dict
        header_color: override header background color
    """
    from templates.styles import (
        CONTENT_LEFT, CONTENT_TOP, get_colors, get_font_name
    )

    left = left or CONTENT_LEFT
    top = top or CONTENT_TOP
    n_rows = len(rows)
    n_cols = len(rows[0]) if rows else 0

    if n_rows == 0 or n_cols == 0:
        return None

    total_width = sum(col_widths) if col_widths else Inches(12)
    row_height = Inches(0.35)
    total_height = height or Inches(row_height.inches * n_rows)

    table_shape = slide.shapes.add_table(
        n_rows, n_cols, left, top, total_width, total_height
    )
    table = table_shape.table

    colors = get_colors(config)
    font_name = get_font_name(config)
    hdr_color = header_color or colors["primary"]

    # Set column widths
    for i, w in enumerate(col_widths):
        if i < n_cols:
            table.columns[i].width = w

    # Populate and format
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(cell_text) if cell_text is not None else ""

            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = font_name
                paragraph.font.size = Pt(9)

                if row_idx == 0:
                    # Header row
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = colors["text_light"]
                    paragraph.font.size = Pt(10)
                else:
                    paragraph.font.color.rgb = colors["text_dark"]

            # Header background
            if row_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = hdr_color

            # Alternating row colors
            elif row_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = hex_to_rgb("F0F4F8")

    return table_shape


def add_kpi_box(slide, label, value, config,
                left=Inches(0.5), top=Inches(1.5),
                width=Inches(2.5), height=Inches(1.2),
                value_color=None):
    """Add a KPI metric box (label + large value)."""
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from templates.styles import get_colors, get_font_name

    colors = get_colors(config)
    font_name = get_font_name(config)

    # Background shape
    shape = slide.shapes.add_shape(
        1,  # RECTANGLE
        left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb("F0F4F8")
    shape.line.color.rgb = hex_to_rgb("DDDDDD")
    shape.line.width = Pt(0.5)

    # Label
    txBox = slide.shapes.add_textbox(
        left + Inches(0.1), top + Inches(0.1),
        width - Inches(0.2), Inches(0.3)
    )
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(9)
    p.font.color.rgb = hex_to_rgb("666666")
    p.font.name = font_name

    # Value
    txBox2 = slide.shapes.add_textbox(
        left + Inches(0.1), top + Inches(0.4),
        width - Inches(0.2), Inches(0.6)
    )
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = str(value)
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = value_color or colors["primary"]
    p2.font.name = font_name

    return shape


def add_bullet_list(slide, items, config,
                    left=None, top=None, width=None, height=None,
                    font_size_pt=11):
    """Add a bullet-point list to a slide."""
    from templates.styles import (
        CONTENT_LEFT, CONTENT_TOP, CONTENT_WIDTH, CONTENT_HEIGHT,
        get_colors, get_font_name
    )

    left = left or CONTENT_LEFT
    top = top or CONTENT_TOP
    width = width or CONTENT_WIDTH
    height = height or CONTENT_HEIGHT

    colors = get_colors(config)
    font_name = get_font_name(config)

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(font_size_pt)
        p.font.color.rgb = colors["text_dark"]
        p.font.name = font_name
        p.space_after = Pt(6)

    return txBox
