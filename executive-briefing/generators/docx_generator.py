"""
Word Document Generator — TTEC Executive Intelligence Briefing
Generates a detailed narrative appendix with 10 sections.
"""

import sys
import os
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from templates.styles import format_number


# ---------------------------------------------------------------------------
# Shared formatting helpers
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x00, 0x33, 0x66)
TEAL = RGBColor(0x00, 0x7A, 0x8C)
DARK_TEXT = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x99, 0x99, 0x99)


def _add_heading(doc, text, level=1):
    """Add a styled heading."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = NAVY
    return heading


def _add_body(doc, text):
    """Add a body paragraph."""
    p = doc.add_paragraph(text)
    p.style.font.name = "Calibri"
    p.style.font.size = Pt(11)
    return p


def _add_bullet(doc, text):
    """Add a bullet point."""
    p = doc.add_paragraph(text, style="List Bullet")
    return p


def _add_table(doc, headers, rows):
    """Add a formatted table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Headers
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)

    # Data rows
    for row_idx, row_data in enumerate(rows):
        for col_idx, value in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = str(value) if value is not None else "N/A"
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    return table


# ---------------------------------------------------------------------------
# Section 1: Industry Analysis
# ---------------------------------------------------------------------------
def _section_industry(doc, data):
    _add_heading(doc, "1. Industry Analysis")

    ind = data.get("industry", {})
    market = ind.get("market_size", {})

    _add_body(doc, (
        f"The global CX/BPO market was valued at ${market.get('current_value_billions', 102)} billion "
        f"in {market.get('current_year', 2024)} and is projected to reach "
        f"${market.get('projected_value_billions', 296)} billion by {market.get('projected_year', 2033)}, "
        f"representing a {market.get('cagr_pct', 12.8)}% CAGR. "
        f"This growth is being driven by enterprise demand for omnichannel customer engagement, "
        f"digital transformation initiatives, and the rapid integration of AI and automation technologies."
    ))

    _add_heading(doc, "Growth Drivers", level=2)
    for driver in ind.get("key_growth_drivers", []):
        _add_bullet(doc, driver)

    _add_heading(doc, "Key Headwinds", level=2)
    for headwind in ind.get("key_headwinds", []):
        _add_bullet(doc, headwind)

    _add_heading(doc, "Geographic Distribution", level=2)
    geos = ind.get("geographic_distribution", [])
    if geos:
        headers = ["Region", "Market Share", "Notes"]
        rows = [[g.get("region", ""), f"{g.get('share_pct', '')}%", g.get("notes", "")] for g in geos]
        _add_table(doc, headers, rows)

    _add_heading(doc, "Industry Concentration", level=2)
    conc = ind.get("industry_concentration", {})
    _add_body(doc, (
        f"The top 5 providers hold approximately {conc.get('top_5_share_pct', 25)}% of the market, "
        f"while the top 10 account for {conc.get('top_10_share_pct', 35)}%. "
        f"The industry is {conc.get('trend', 'consolidating through M&A')}."
    ))


# ---------------------------------------------------------------------------
# Section 2: Technology Landscape
# ---------------------------------------------------------------------------
def _section_technology(doc, data):
    _add_heading(doc, "2. Technology Landscape")

    tech = data.get("technology_trends", {})
    ai = tech.get("ai_adoption", {})

    _add_body(doc, (
        f"{ai.get('headline', '')}. The integration of artificial intelligence into customer experience "
        f"operations represents the most significant transformation in the BPO industry's history."
    ))

    _add_heading(doc, "AI Adoption Metrics", level=2)
    stats = ai.get("stats", [])
    for stat in stats:
        val = stat.get("value", "")
        unit = stat.get("unit", "")
        if unit == "percent":
            display = f"{val}%"
        elif unit == "billions_usd":
            display = f"${val}B"
        else:
            display = str(val)
        _add_bullet(doc, f"{stat.get('metric', '')}: {display} (Source: {stat.get('source', '')})")

    _add_heading(doc, "Cloud & Digital CX Trends", level=2)
    cloud = tech.get("cloud_and_digital", {})
    trends = cloud.get("trends", [])
    if trends:
        headers = ["Trend", "Adoption (%)", "Growth Rate (%)", "Description"]
        rows = [
            [t.get("name", ""), f"{t.get('adoption_pct', '')}%",
             f"{t.get('growth_rate_pct', '')}%", t.get("description", "")]
            for t in trends
        ]
        _add_table(doc, headers, rows)

    _add_heading(doc, "Emerging Business Models", level=2)
    models = tech.get("emerging_models", [])
    for m in models:
        _add_heading(doc, m.get("name", ""), level=3)
        _add_body(doc, f"{m.get('description', '')} (Stage: {m.get('adoption_stage', '')})")
        _add_body(doc, f"Impact: {m.get('impact', '')}")


# ---------------------------------------------------------------------------
# Section 3: AI Disruption Deep Dive
# ---------------------------------------------------------------------------
def _section_ai_disruption(doc, data):
    _add_heading(doc, "3. AI Disruption Deep Dive")

    tech = data.get("technology_trends", {})
    assessment = tech.get("ai_disruption_assessment", {})
    scenarios = assessment.get("tam_compression_scenario", {})

    _add_body(doc, (
        "The rise of generative AI and agentic AI systems poses a fundamental challenge to the traditional "
        "headcount-based CX/BPO business model. This section analyzes three scenarios for TAM compression "
        "and assesses TTEC's positioning relative to peers."
    ))

    _add_heading(doc, "TAM Compression Scenarios", level=2)
    if scenarios:
        headers = ["Scenario", "AI Adoption", "Revenue Impact", "Timeline", "Description"]
        rows = []
        for key in ["scenario_low", "scenario_mid", "scenario_high"]:
            s = scenarios.get(key, {})
            rows.append([
                key.replace("scenario_", "").capitalize(),
                f"{s.get('ai_adoption_pct', '')}%",
                f"{s.get('revenue_impact_pct', '')}%",
                s.get("timeline", ""),
                s.get("description", ""),
            ])
        _add_table(doc, headers, rows)

    readiness = assessment.get("ttec_readiness", {})

    _add_heading(doc, "TTEC AI Readiness — Strengths", level=2)
    for s in readiness.get("strengths", []):
        _add_bullet(doc, s)

    _add_heading(doc, "TTEC AI Readiness — Vulnerabilities", level=2)
    for v in readiness.get("vulnerabilities", []):
        _add_bullet(doc, v)

    _add_heading(doc, "Peer AI Readiness Comparison", level=2)
    peer_comp = readiness.get("peer_comparison", [])
    if peer_comp:
        headers = ["Company", "AI Readiness", "Notes"]
        rows = [[p.get("company", ""), p.get("ai_readiness", ""), p.get("notes", "")] for p in peer_comp]
        _add_table(doc, headers, rows)


# ---------------------------------------------------------------------------
# Section 4: Competitive Profiles
# ---------------------------------------------------------------------------
def _section_competitive_profiles(doc, data):
    _add_heading(doc, "4. Competitive Profiles")

    _add_body(doc, (
        "The following profiles cover TTEC's 12 key competitors across the CX/BPO landscape, "
        "organized by tier. Tier 1 represents large-cap direct competitors; Tier 2 includes "
        "mid-cap public peers and private competitors."
    ))

    comps = data.get("competitors", {}).get("competitors", [])
    for c in comps:
        name = c.get("name", "")
        ticker = c.get("ticker")
        header = f"{name} ({ticker})" if ticker else name
        _add_heading(doc, header, level=2)

        # Key facts
        facts = []
        if c.get("headquarters"):
            facts.append(f"HQ: {c['headquarters']}")
        if c.get("employees"):
            facts.append(f"Employees: {c['employees']:,}")
        if c.get("revenue_millions"):
            facts.append(f"Revenue: {format_number(c['revenue_millions'], 'money')}")
        if c.get("adj_ebitda_margin_pct"):
            facts.append(f"EBITDA Margin: {c['adj_ebitda_margin_pct']}%")
        if c.get("ev_ebitda_multiple"):
            facts.append(f"EV/EBITDA: {c['ev_ebitda_multiple']}x")
        if facts:
            _add_body(doc, " | ".join(facts))

        # Strategy
        if c.get("strategy"):
            _add_body(doc, f"Strategy: {c['strategy']}")

        # Recent highlights
        highlights = c.get("recent_highlights", [])
        if highlights:
            _add_heading(doc, "Recent Highlights", level=3)
            for h in highlights:
                _add_bullet(doc, h)

        # Key verticals
        verticals = c.get("key_verticals", [])
        if verticals:
            _add_body(doc, f"Key Verticals: {', '.join(verticals)}")

        doc.add_paragraph("")  # Spacing


# ---------------------------------------------------------------------------
# Section 5: Financial Analysis
# ---------------------------------------------------------------------------
def _section_financial_analysis(doc, data):
    _add_heading(doc, "5. Financial Analysis")

    _add_body(doc, (
        "This section provides a detailed comparison of TTEC's financial performance "
        "against its peer group, including revenue trends, profitability metrics, and "
        "valuation multiples."
    ))

    _add_heading(doc, "TTEC Historical Performance", level=2)
    co = data.get("company_overview", {})
    annual = co.get("annual_results", [])
    if annual:
        headers = ["Year", "Revenue ($M)", "EBITDA ($M)", "Margin (%)", "FCF ($M)", "Net Debt ($M)"]
        rows = [
            [str(yr.get("year")),
             str(yr.get("revenue_millions", "")),
             str(yr.get("adj_ebitda_millions", "")),
             f"{yr.get('adj_ebitda_margin_pct', '')}%",
             str(yr.get("free_cash_flow_millions", "")),
             str(yr.get("net_debt_millions", ""))]
            for yr in annual
        ]
        _add_table(doc, headers, rows)

    _add_heading(doc, "Peer Valuation Summary", level=2)
    comps = data.get("competitors", {}).get("competitors", [])
    public_comps = [c for c in comps if c.get("public") and c.get("ev_ebitda_multiple")]
    if public_comps:
        headers = ["Company", "Revenue ($M)", "EBITDA Margin (%)", "EV/EBITDA", "Market Cap ($M)"]
        rows = [
            [c.get("name", ""),
             str(c.get("revenue_millions", "")),
             f"{c.get('adj_ebitda_margin_pct', '')}%",
             f"{c.get('ev_ebitda_multiple', '')}x",
             str(c.get("market_cap_millions", ""))]
            for c in public_comps
        ]
        _add_table(doc, headers, rows)

    _add_heading(doc, "FY2026 Guidance", level=2)
    guidance = co.get("guidance_2026", {})
    if guidance:
        _add_body(doc, (
            f"Revenue: ${guidance.get('revenue_low_millions', '')}-${guidance.get('revenue_high_millions', '')}M | "
            f"EBITDA: ${guidance.get('adj_ebitda_low_millions', '')}-${guidance.get('adj_ebitda_high_millions', '')}M"
        ))
        for c in guidance.get("commentary", []):
            _add_bullet(doc, c)


# ---------------------------------------------------------------------------
# Section 6: Analyst & Market Sentiment
# ---------------------------------------------------------------------------
def _section_analyst_sentiment(doc, data):
    _add_heading(doc, "6. Analyst & Market Sentiment")

    analyst = data.get("analyst_consensus", {})
    ttec = analyst.get("ttec", {})

    _add_body(doc, (
        f"TTEC currently carries a consensus rating of '{ttec.get('consensus_rating', 'N/A')}' "
        f"from {ttec.get('num_analysts', 'N/A')} sell-side analysts. "
        f"The median price target is ${ttec.get('price_target', {}).get('median', 'N/A')}, "
        f"compared to a current stock price of approximately ${ttec.get('current_price', 'N/A')}."
    ))

    _add_heading(doc, "TTEC Analyst Sentiment", level=2)
    for note in ttec.get("sentiment_notes", []):
        _add_bullet(doc, note)

    _add_heading(doc, "Sector Themes", level=2)
    for theme in analyst.get("sector_themes", []):
        _add_bullet(doc, theme)

    _add_heading(doc, "Peer Consensus Comparison", level=2)
    peers = analyst.get("key_peers", [])
    if peers:
        headers = ["Company", "Rating", "Price Target", "Current Price", "Fwd EV/EBITDA"]
        rows = [
            [p.get("name", ""), p.get("consensus_rating", ""),
             f"${p.get('price_target_median', '')}",
             f"${p.get('current_price', '')}",
             f"{p.get('fy_forward_ev_ebitda', '')}x"]
            for p in peers
        ]
        _add_table(doc, headers, rows)


# ---------------------------------------------------------------------------
# Section 7: Capital Allocation Review
# ---------------------------------------------------------------------------
def _section_capital_allocation(doc, data):
    _add_heading(doc, "7. Capital Allocation Review")

    cap = data.get("capital_allocation", {})
    ttec = cap.get("ttec", {})
    bs = ttec.get("balance_sheet", {})

    _add_body(doc, (
        f"TTEC ended FY2025 with net debt of ${bs.get('net_debt_millions', '')}M "
        f"and a net leverage ratio of {bs.get('net_leverage_ratio', '')}x, "
        f"down from {bs.get('prior_year_leverage_ratio', '')}x at the end of FY2024. "
        f"The company reduced credit facility borrowings by ${bs.get('debt_reduction_millions', '')}M year-over-year."
    ))

    _add_heading(doc, "M&A Capacity", level=2)
    ma_cap = ttec.get("ma_capacity", {})
    _add_body(doc, (
        f"Estimated acquisition capacity: ${ma_cap.get('estimated_acquisition_capacity_millions', '')}M. "
        f"{ma_cap.get('notes', '')}"
    ))

    _add_heading(doc, "Peer Capital Deployment", level=2)
    peers = cap.get("peer_capital_allocation", [])
    if peers:
        headers = ["Company", "Leverage", "Debt Paydown", "M&A", "Buybacks", "Strategy"]
        rows = [
            [p.get("name", ""),
             f"{p.get('net_leverage_ratio', '')}x",
             f"${p.get('fy_debt_reduction_millions', '')}M",
             f"${p.get('fy_acquisitions_millions', '')}M",
             f"${p.get('fy_buybacks_millions', '')}M",
             p.get("strategy", "")]
            for p in peers
        ]
        _add_table(doc, headers, rows)

    _add_heading(doc, "Key Observations", level=2)
    for obs in cap.get("key_observations", []):
        _add_bullet(doc, obs)


# ---------------------------------------------------------------------------
# Section 8: Labor Market Analysis
# ---------------------------------------------------------------------------
def _section_labor_market(doc, data):
    _add_heading(doc, "8. Labor Market Analysis")

    labor = data.get("labor_analytics", {})

    total = labor.get("industry_workforce", {}).get("total_global_cx_bpo_employees")
    _add_body(doc, (
        f"The global CX/BPO industry employs an estimated {total:,} workers "
        f"across key delivery geographies. Labor costs represent 65-75% of total "
        f"operating expenses for most CX providers, making workforce dynamics "
        f"a critical factor in industry economics."
    ))

    _add_heading(doc, "Delivery Geography Analysis", level=2)
    geos = labor.get("industry_workforce", {}).get("top_delivery_geographies", [])
    if geos:
        headers = ["Country", "Employees", "Avg Wage", "Wage Growth", "Attrition", "Trend"]
        rows = [
            [g.get("country", ""),
             f"{g.get('employees_estimate', 0):,}",
             f"${g.get('avg_annual_wage_usd', 0):,}",
             f"{g.get('wage_growth_yoy_pct', '')}%",
             f"{g.get('attrition_rate_pct', '')}%",
             g.get("trend", "")]
            for g in geos
        ]
        _add_table(doc, headers, rows)

    _add_heading(doc, "AI Impact on Workforce", level=2)
    ai_impact = labor.get("ai_impact_on_workforce", {})
    _add_body(doc, ai_impact.get("headline", ""))
    for effect in ai_impact.get("effects", []):
        _add_bullet(doc, effect)

    _add_heading(doc, "Key Observations", level=2)
    for obs in labor.get("key_observations", []):
        _add_bullet(doc, obs)


# ---------------------------------------------------------------------------
# Section 9: Strategic Assessment
# ---------------------------------------------------------------------------
def _section_strategic_assessment(doc, data):
    _add_heading(doc, "9. Strategic Assessment")

    co = data.get("company_overview", {})
    cap = data.get("capital_allocation", {})

    _add_body(doc, (
        "This section provides a SWOT-style assessment of TTEC's strategic positioning "
        "based on the competitive, financial, and market analysis presented in this briefing."
    ))

    _add_heading(doc, "Strengths", level=2)
    strengths = [
        "Dual-segment model (Digital + Engage) provides technology and services synergies",
        "TTEC Digital growing 9.2% YoY — strong positioning in CX technology consulting",
        "Demonstrated ability to generate free cash flow ($83M in FY2025)",
        "Early AI adoption strategy with near 100% client AI adoption target",
        "20-country delivery footprint provides geographic flexibility",
    ]
    for s in strengths:
        _add_bullet(doc, s)

    _add_heading(doc, "Weaknesses", level=2)
    weaknesses = [
        "Highest leverage in peer group (3.58x) limits strategic flexibility",
        "Revenue declining (-3.3% FY2025) as underperforming contracts are rationalized",
        "TTEC Engage still primarily headcount-based, vulnerable to AI disruption",
        "Smaller scale vs. Concentrix and Teleperformance limits R&D investment capacity",
        "$205.4M goodwill impairment signals Digital segment challenges",
    ]
    for w in weaknesses:
        _add_bullet(doc, w)

    _add_heading(doc, "Opportunities", level=2)
    opportunities = [
        "AI transformation consulting demand growing rapidly — aligns with Digital segment",
        "Compressed peer valuations may create attractive M&A opportunities (once deleveraged)",
        "Outcome-based pricing models could increase margins if AI execution succeeds",
        "Nearshore expansion (LatAm) can improve cost structure and client proximity",
        "Enterprise complexity in AI deployment creates demand for experienced integrators",
    ]
    for o in opportunities:
        _add_bullet(doc, o)

    _add_heading(doc, "Threats", level=2)
    threats = [
        "Aggressive AI adoption could compress CX TAM by 18-30% over 3-5 years",
        "Well-capitalized competitors (TaskUs, ExlService) investing heavily in AI",
        "Macro downturn could accelerate contract rationalization by clients",
        "Wage inflation in key delivery markets narrowing offshore cost advantage",
        "Regulatory changes (EU AI Act, data privacy) increasing compliance costs",
    ]
    for t in threats:
        _add_bullet(doc, t)


# ---------------------------------------------------------------------------
# Section 10: Sources & Methodology
# ---------------------------------------------------------------------------
def _section_sources(doc, data):
    _add_heading(doc, "10. Sources & Methodology")

    _add_heading(doc, "Primary Data Sources", level=2)
    sources = [
        "TTEC Holdings Investor Relations (investors.ttec.com) — Earnings releases, SEC filings, presentations",
        "SEC/EDGAR — 10-K, 10-Q, and proxy filings for all public peer companies",
        "FactSet — Financial data validation, analyst consensus, valuation multiples",
        "Grand View Research — CX BPO market sizing and growth projections",
        "Gartner — Customer service technology trends and forecasts",
        "Everest Group — CX/BPO industry assessments and provider rankings",
        "ISG — Outsourcing industry analysis and provider evaluations",
        "Bureau of Labor Statistics — US employment and wage data",
        "Company IR websites — Peer company earnings releases and investor presentations",
        "Verint — Contact center AI adoption trends",
        "Ryan Strategic Advisory — CX/BPO M&A activity tracker",
    ]
    for s in sources:
        _add_bullet(doc, s)

    _add_heading(doc, "Methodology Notes", level=2)
    _add_bullet(doc, "Financial data for public companies sourced from most recent annual filings")
    _add_bullet(doc, "Private company data (Alorica, Foundever, Sutherland) estimated from industry reports and press releases")
    _add_bullet(doc, "Valuation multiples based on latest available market data")
    _add_bullet(doc, "AI adoption statistics from multiple industry surveys — figures represent industry averages")
    _add_bullet(doc, "Market sizing data from Grand View Research; alternative estimates may vary by 10-15%")

    _add_heading(doc, "Disclaimer", level=2)
    _add_body(doc, (
        "This briefing is prepared for internal use by TTEC Holdings' executive leadership "
        "and Board of Directors. It contains estimates, projections, and forward-looking statements "
        "based on publicly available information. Financial data for private companies is estimated. "
        "All data should be validated against primary sources before use in investment or strategic decisions. "
        "This document is classified as CONFIDENTIAL."
    ))


# =========================================================================
# Main generation function
# =========================================================================
def generate(data: dict, config: dict, output_path: str):
    """Generate the full Word document."""
    doc = Document()

    # Title
    title = doc.add_heading(
        f"TTEC Holdings — Executive Intelligence Briefing",
        level=0,
    )
    for run in title.runs:
        run.font.color.rgb = NAVY

    quarter = config.get("quarter", "Q1")
    year = config.get("year", 2026)
    subtitle = doc.add_paragraph(f"{quarter} {year} — Detailed Analysis & Supporting Documentation")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(14)
        run.font.color.rgb = TEAL

    classification = doc.add_paragraph(config.get("classification", "CONFIDENTIAL"))
    classification.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in classification.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = GRAY
        run.bold = True

    doc.add_page_break()

    # Table of Contents placeholder
    _add_heading(doc, "Table of Contents")
    toc_items = [
        "1. Industry Analysis",
        "2. Technology Landscape",
        "3. AI Disruption Deep Dive",
        "4. Competitive Profiles (12 companies)",
        "5. Financial Analysis",
        "6. Analyst & Market Sentiment",
        "7. Capital Allocation Review",
        "8. Labor Market Analysis",
        "9. Strategic Assessment (SWOT)",
        "10. Sources & Methodology",
    ]
    for item in toc_items:
        _add_body(doc, item)

    doc.add_page_break()

    # Generate all sections
    _section_industry(doc, data)
    doc.add_page_break()

    _section_technology(doc, data)
    doc.add_page_break()

    _section_ai_disruption(doc, data)
    doc.add_page_break()

    _section_competitive_profiles(doc, data)
    doc.add_page_break()

    _section_financial_analysis(doc, data)
    doc.add_page_break()

    _section_analyst_sentiment(doc, data)
    doc.add_page_break()

    _section_capital_allocation(doc, data)
    doc.add_page_break()

    _section_labor_market(doc, data)
    doc.add_page_break()

    _section_strategic_assessment(doc, data)
    doc.add_page_break()

    _section_sources(doc, data)

    doc.save(output_path)
    print(f"  Word document saved: {output_path}")
