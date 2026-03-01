"""
Excel Workbook Generator — TTEC Executive Intelligence Briefing
Generates an 8-tab workbook with financial data, benchmarking, and analysis.
"""

import sys
import os
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

from templates.styles import hex_to_rgb, format_number


# ---------------------------------------------------------------------------
# Shared formatting helpers
# ---------------------------------------------------------------------------
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
ALT_ROW_FILL = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
BODY_FONT = Font(name="Calibri", size=10, color="333333")
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="003366")
THIN_BORDER = Border(
    left=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"),
    bottom=Side(style="thin", color="DDDDDD"),
)


def _setup_sheet(ws, title, columns, col_widths):
    """Apply standard formatting to a worksheet."""
    ws.title = title
    # Column headers
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    # Column widths
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze header row
    ws.freeze_panes = "A2"


def _write_row(ws, row_idx, values, is_alt=False):
    """Write a row of data with formatting."""
    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = BODY_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center")
        if is_alt:
            cell.fill = ALT_ROW_FILL
        # Format numbers
        if isinstance(value, float):
            if abs(value) < 100:  # Likely a percentage or ratio
                cell.number_format = "0.0"
            else:
                cell.number_format = "#,##0"
        elif isinstance(value, int) and abs(value) > 100:
            cell.number_format = "#,##0"


# ---------------------------------------------------------------------------
# Tab 1: TTEC Financials
# ---------------------------------------------------------------------------
def _tab_ttec_financials(wb, data):
    ws = wb.active
    columns = [
        "Year", "Revenue ($M)", "Adj. EBITDA ($M)", "EBITDA Margin (%)",
        "Operating Income ($M)", "Net Income ($M)", "Free Cash Flow ($M)",
        "Net Debt ($M)", "Employees",
    ]
    col_widths = [10, 16, 16, 16, 18, 16, 18, 14, 14]
    _setup_sheet(ws, "TTEC Financials", columns, col_widths)

    co = data.get("company_overview", {})
    annual = co.get("annual_results", [])

    for i, yr in enumerate(annual):
        row_idx = i + 2
        values = [
            yr.get("year"),
            yr.get("revenue_millions"),
            yr.get("adj_ebitda_millions"),
            yr.get("adj_ebitda_margin_pct"),
            yr.get("operating_income_millions"),
            yr.get("net_income_millions"),
            yr.get("free_cash_flow_millions"),
            yr.get("net_debt_millions"),
            yr.get("employees"),
        ]
        _write_row(ws, row_idx, values, is_alt=(i % 2 == 1))

    # Add guidance row
    guidance = co.get("guidance_2026", {})
    if guidance:
        row_idx = len(annual) + 2
        rev_mid = (guidance.get("revenue_low_millions", 0) + guidance.get("revenue_high_millions", 0)) / 2
        ebitda_mid = (guidance.get("adj_ebitda_low_millions", 0) + guidance.get("adj_ebitda_high_millions", 0)) / 2
        values = [
            "2026E",
            rev_mid,
            ebitda_mid,
            ebitda_mid / rev_mid * 100 if rev_mid else 0,
            None, None, None, None, None,
        ]
        _write_row(ws, row_idx, values, is_alt=(len(annual) % 2 == 1))
        # Highlight guidance row
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10, color="003366", italic=True)

    # Segment detail section
    row_idx = len(annual) + 4
    ws.cell(row=row_idx, column=1, value="Segment Detail (Latest Quarter)").font = TITLE_FONT
    row_idx += 1

    seg_cols = ["Segment", "Revenue ($M)", "Revenue Growth (%)", "Non-GAAP Op. Margin (%)"]
    for col_idx, col_name in enumerate(seg_cols, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

    qs = co.get("quarterly_segments", {})
    for seg_key in ["ttec_digital", "ttec_engage"]:
        seg = qs.get(seg_key, {})
        if seg:
            row_idx += 1
            values = [
                seg_key.replace("_", " ").title(),
                seg.get("revenue_millions"),
                seg.get("revenue_growth_yoy_pct"),
                seg.get("non_gaap_operating_margin_pct"),
            ]
            _write_row(ws, row_idx, values)


# ---------------------------------------------------------------------------
# Tab 2: Peer Benchmarking
# ---------------------------------------------------------------------------
def _tab_peer_benchmarking(wb, data):
    ws = wb.create_sheet()
    columns = [
        "Company", "Ticker", "Revenue ($M)", "Revenue Growth (%)",
        "EBITDA ($M)", "EBITDA Margin (%)", "Employees",
        "Rev/Employee ($K)", "Tier", "Public/Private",
    ]
    col_widths = [22, 10, 14, 16, 14, 16, 14, 16, 8, 14]
    _setup_sheet(ws, "Peer Benchmarking", columns, col_widths)

    # TTEC first
    co = data.get("company_overview", {})
    annual = co.get("annual_results", [])
    if annual:
        latest = annual[-1]
        rev = latest.get("revenue_millions", 0)
        emp = latest.get("employees", 1)
        values = [
            "TTEC Holdings", "TTEC",
            rev,
            None,  # Growth not easily calculated here
            latest.get("adj_ebitda_millions"),
            latest.get("adj_ebitda_margin_pct"),
            emp,
            round(rev / emp * 1000, 1) if emp else None,
            "—", "Public",
        ]
        _write_row(ws, 2, values)
        # Highlight TTEC row
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=2, column=col_idx)
            cell.font = Font(name="Calibri", size=10, bold=True, color="003366")

    comps = data.get("competitors", {}).get("competitors", [])
    for i, c in enumerate(comps):
        row_idx = i + 3
        rev = c.get("revenue_millions", 0) or 0
        emp = c.get("employees", 1) or 1
        values = [
            c.get("name"),
            c.get("ticker", "—"),
            rev,
            c.get("revenue_growth_yoy_pct"),
            c.get("adj_ebitda_millions"),
            c.get("adj_ebitda_margin_pct"),
            c.get("employees"),
            round(rev / emp * 1000, 1) if emp else None,
            f"Tier {c.get('tier', '')}",
            "Public" if c.get("public") else "Private",
        ]
        _write_row(ws, row_idx, values, is_alt=(i % 2 == 0))


# ---------------------------------------------------------------------------
# Tab 3: Valuation Comps
# ---------------------------------------------------------------------------
def _tab_valuation_comps(wb, data):
    ws = wb.create_sheet()
    columns = [
        "Company", "Ticker", "Market Cap ($M)", "EV/EBITDA",
        "Revenue ($M)", "EBITDA Margin (%)",
    ]
    col_widths = [22, 10, 16, 14, 14, 16]
    _setup_sheet(ws, "Valuation Comps", columns, col_widths)

    comps = data.get("competitors", {}).get("competitors", [])
    public_comps = [c for c in comps if c.get("public")]

    for i, c in enumerate(public_comps):
        row_idx = i + 2
        values = [
            c.get("name"),
            c.get("ticker"),
            c.get("market_cap_millions"),
            c.get("ev_ebitda_multiple"),
            c.get("revenue_millions"),
            c.get("adj_ebitda_margin_pct"),
        ]
        _write_row(ws, row_idx, values, is_alt=(i % 2 == 1))


# ---------------------------------------------------------------------------
# Tab 4: Analyst Consensus
# ---------------------------------------------------------------------------
def _tab_analyst_consensus(wb, data):
    ws = wb.create_sheet()
    columns = [
        "Company", "Ticker", "Rating", "# Analysts",
        "Price Target ($)", "Current Price ($)", "Fwd EV/EBITDA",
    ]
    col_widths = [22, 10, 12, 12, 16, 16, 14]
    _setup_sheet(ws, "Analyst Consensus", columns, col_widths)

    analyst = data.get("analyst_consensus", {})

    # TTEC
    ttec = analyst.get("ttec", {})
    pt = ttec.get("price_target", {})
    values = [
        "TTEC Holdings", "TTEC",
        ttec.get("consensus_rating"),
        ttec.get("num_analysts"),
        pt.get("median"),
        ttec.get("current_price"),
        None,
    ]
    _write_row(ws, 2, values)
    for col_idx in range(1, len(columns) + 1):
        ws.cell(row=2, column=col_idx).font = Font(name="Calibri", size=10, bold=True, color="003366")

    # Peers
    peers = analyst.get("key_peers", [])
    for i, p in enumerate(peers):
        row_idx = i + 3
        values = [
            p.get("name"),
            p.get("ticker"),
            p.get("consensus_rating"),
            p.get("num_analysts"),
            p.get("price_target_median"),
            p.get("current_price"),
            p.get("fy_forward_ev_ebitda"),
        ]
        _write_row(ws, row_idx, values, is_alt=(i % 2 == 0))


# ---------------------------------------------------------------------------
# Tab 5: Capital Allocation
# ---------------------------------------------------------------------------
def _tab_capital_allocation(wb, data):
    ws = wb.create_sheet()
    columns = [
        "Company", "Net Leverage", "Debt Reduction ($M)",
        "Acquisitions ($M)", "Buybacks ($M)", "Div. Yield (%)", "Strategy",
    ]
    col_widths = [22, 14, 18, 16, 14, 14, 40]
    _setup_sheet(ws, "Capital Allocation", columns, col_widths)

    cap = data.get("capital_allocation", {})

    # TTEC
    ttec = cap.get("ttec", {})
    bs = ttec.get("balance_sheet", {})
    deploy = ttec.get("capital_deployment_fy2025", {})
    values = [
        "TTEC Holdings",
        bs.get("net_leverage_ratio"),
        deploy.get("debt_reduction_millions"),
        deploy.get("acquisitions_millions"),
        deploy.get("buybacks_millions"),
        ttec.get("dividend", {}).get("yield_pct"),
        "Deleveraging priority, limited M&A capacity",
    ]
    _write_row(ws, 2, values)
    for col_idx in range(1, len(columns) + 1):
        ws.cell(row=2, column=col_idx).font = Font(name="Calibri", size=10, bold=True, color="003366")

    # Peers
    peers = cap.get("peer_capital_allocation", [])
    for i, p in enumerate(peers):
        row_idx = i + 3
        values = [
            p.get("name"),
            p.get("net_leverage_ratio"),
            p.get("fy_debt_reduction_millions"),
            p.get("fy_acquisitions_millions"),
            p.get("fy_buybacks_millions"),
            p.get("dividend_yield_pct"),
            p.get("strategy", ""),
        ]
        _write_row(ws, row_idx, values, is_alt=(i % 2 == 0))


# ---------------------------------------------------------------------------
# Tab 6: Labor Analytics
# ---------------------------------------------------------------------------
def _tab_labor_analytics(wb, data):
    ws = wb.create_sheet()
    columns = [
        "Country", "Est. Employees", "Avg Annual Wage (USD)",
        "Wage Growth YoY (%)", "Attrition Rate (%)", "Trend",
    ]
    col_widths = [18, 16, 20, 18, 16, 45]
    _setup_sheet(ws, "Labor Analytics", columns, col_widths)

    labor = data.get("labor_analytics", {})
    geos = labor.get("industry_workforce", {}).get("top_delivery_geographies", [])

    for i, g in enumerate(geos):
        row_idx = i + 2
        values = [
            g.get("country"),
            g.get("employees_estimate"),
            g.get("avg_annual_wage_usd"),
            g.get("wage_growth_yoy_pct"),
            g.get("attrition_rate_pct"),
            g.get("trend", ""),
        ]
        _write_row(ws, row_idx, values, is_alt=(i % 2 == 1))

    # Delivery mix section
    row_idx = len(geos) + 3
    ws.cell(row=row_idx, column=1, value="Delivery Mix Trends").font = TITLE_FONT
    row_idx += 1

    mix_cols = ["Model", "Current Share (%)", "Projected 2028 (%)", "Trend"]
    for col_idx, col_name in enumerate(mix_cols, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

    mix = labor.get("delivery_mix_trends", [])
    for i, m in enumerate(mix):
        row_idx += 1
        values = [
            m.get("model"),
            m.get("current_share_pct"),
            m.get("projected_2028_share_pct"),
            m.get("trend", ""),
        ]
        _write_row(ws, row_idx, values, is_alt=(i % 2 == 1))


# ---------------------------------------------------------------------------
# Tab 7: M&A Tracker
# ---------------------------------------------------------------------------
def _tab_ma_tracker(wb, data):
    ws = wb.create_sheet()
    columns = [
        "Date", "Acquirer", "Target", "Deal Value ($M)",
        "EV/Revenue", "EV/EBITDA", "Rationale", "Status",
    ]
    col_widths = [12, 18, 22, 16, 12, 12, 45, 16]
    _setup_sheet(ws, "M&A Tracker", columns, col_widths)

    ma = data.get("ma_transactions", {})
    deals = ma.get("recent_transactions", [])

    for i, d in enumerate(deals):
        row_idx = i + 2
        values = [
            d.get("date"),
            d.get("acquirer"),
            d.get("target"),
            d.get("deal_value_millions"),
            d.get("ev_revenue_multiple"),
            d.get("ev_ebitda_multiple"),
            d.get("rationale", ""),
            d.get("status", ""),
        ]
        _write_row(ws, row_idx, values, is_alt=(i % 2 == 1))

    # Valuation benchmarks section
    benchmarks = ma.get("valuation_benchmarks", {})
    if benchmarks:
        row_idx = len(deals) + 3
        ws.cell(row=row_idx, column=1, value="M&A Valuation Benchmarks").font = TITLE_FONT
        row_idx += 1

        bench_cols = ["Metric", "Low", "Median", "High", "Notes"]
        for col_idx, col_name in enumerate(bench_cols, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=col_name)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER

        for metric_name, metric_data in benchmarks.items():
            if isinstance(metric_data, dict):
                row_idx += 1
                values = [
                    metric_name.replace("_", " ").title(),
                    metric_data.get("low"),
                    metric_data.get("median"),
                    metric_data.get("high"),
                    metric_data.get("notes", ""),
                ]
                _write_row(ws, row_idx, values)


# ---------------------------------------------------------------------------
# Tab 8: Industry Data
# ---------------------------------------------------------------------------
def _tab_industry_data(wb, data):
    ws = wb.create_sheet()
    columns = [
        "Metric", "Value", "Year/Period", "Source",
    ]
    col_widths = [40, 18, 16, 30]
    _setup_sheet(ws, "Industry Data", columns, col_widths)

    ind = data.get("industry", {})
    market = ind.get("market_size", {})
    tech = data.get("technology_trends", {})

    row_idx = 2
    # Market size data
    rows_data = [
        ("CX BPO Market Size", f"${market.get('current_value_billions', '')}B", str(market.get("current_year", "")), market.get("source", "")),
        ("Projected Market Size", f"${market.get('projected_value_billions', '')}B", str(market.get("projected_year", "")), market.get("source", "")),
        ("Market CAGR", f"{market.get('cagr_pct', '')}%", f"{market.get('current_year', '')}-{market.get('projected_year', '')}", market.get("source", "")),
    ]

    # Add year-by-year data
    by_year = market.get("by_year", {})
    for yr, val in sorted(by_year.items()):
        rows_data.append((f"Market Size ({yr})", f"${val}B", str(yr), market.get("source", "")))

    # AI stats
    ai_stats = tech.get("ai_adoption", {}).get("stats", [])
    for stat in ai_stats:
        unit = stat.get("unit", "")
        val = stat.get("value", "")
        if unit == "percent":
            display = f"{val}%"
        elif unit == "billions_usd":
            display = f"${val}B"
        else:
            display = str(val)
        rows_data.append((stat.get("metric", ""), display, "2025-2026", stat.get("source", "")))

    # Geographic distribution
    geos = ind.get("geographic_distribution", [])
    for g in geos:
        rows_data.append((f"Market Share — {g.get('region', '')}", f"{g.get('share_pct', '')}%", "2024", "Industry estimates"))

    # Industry concentration
    conc = ind.get("industry_concentration", {})
    if conc:
        rows_data.append(("Top 5 Market Share", f"{conc.get('top_5_share_pct', '')}%", "2024", "Industry estimates"))
        rows_data.append(("Top 10 Market Share", f"{conc.get('top_10_share_pct', '')}%", "2024", "Industry estimates"))

    for i, (metric, value, period, source) in enumerate(rows_data):
        _write_row(ws, row_idx + i, [metric, value, period, source], is_alt=(i % 2 == 1))


# =========================================================================
# Main generation function
# =========================================================================
def generate(data: dict, config: dict, output_path: str):
    """Generate the full Excel workbook."""
    wb = Workbook()

    _tab_ttec_financials(wb, data)
    _tab_peer_benchmarking(wb, data)
    _tab_valuation_comps(wb, data)
    _tab_analyst_consensus(wb, data)
    _tab_capital_allocation(wb, data)
    _tab_labor_analytics(wb, data)
    _tab_ma_tracker(wb, data)
    _tab_industry_data(wb, data)

    wb.save(output_path)
    print(f"  Excel saved: {output_path}")
