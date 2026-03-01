"""
Shared styling constants and utilities for TTEC Executive Briefing.
All visual formatting (colors, fonts, positioning) is centralized here.
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

# ---------------------------------------------------------------------------
# Slide dimensions (16:9 widescreen)
# ---------------------------------------------------------------------------
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# ---------------------------------------------------------------------------
# Content positioning
# ---------------------------------------------------------------------------
TITLE_LEFT = Inches(0.5)
TITLE_TOP = Inches(0.3)
TITLE_WIDTH = Inches(12.3)
TITLE_HEIGHT = Inches(0.8)

SUBTITLE_TOP = Inches(1.0)

CONTENT_LEFT = Inches(0.5)
CONTENT_TOP = Inches(1.4)
CONTENT_WIDTH = Inches(12.3)
CONTENT_HEIGHT = Inches(5.5)

# Two-column layout
LEFT_COL_LEFT = Inches(0.5)
LEFT_COL_WIDTH = Inches(5.9)
RIGHT_COL_LEFT = Inches(6.9)
RIGHT_COL_WIDTH = Inches(5.9)

# Chart positioning
CHART_LEFT = Inches(0.5)
CHART_TOP = Inches(1.6)
CHART_WIDTH = Inches(12.3)
CHART_HEIGHT = Inches(5.2)

HALF_CHART_WIDTH = Inches(5.9)

# Footer
FOOTER_TOP = Inches(7.0)
FOOTER_HEIGHT = Inches(0.4)


def hex_to_rgb(hex_color: str) -> RGBColor:
    """Convert hex string (without #) to RGBColor."""
    return RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16),
    )


def get_colors(config: dict) -> dict:
    """Return a dict of RGBColor objects from config."""
    colors = config.get("colors", {})
    return {
        "primary": hex_to_rgb(colors.get("primary", "003366")),
        "secondary": hex_to_rgb(colors.get("secondary", "0066CC")),
        "accent": hex_to_rgb(colors.get("accent", "007A8C")),
        "text_dark": hex_to_rgb(colors.get("text_dark", "333333")),
        "text_light": hex_to_rgb(colors.get("text_light", "FFFFFF")),
        "background": hex_to_rgb(colors.get("background", "F5F5F5")),
        "positive": hex_to_rgb(colors.get("positive", "2E8B57")),
        "negative": hex_to_rgb(colors.get("negative", "CC3333")),
    }


def get_chart_colors(config: dict) -> list:
    """Return list of RGBColor objects for chart series."""
    palette = config.get("colors", {}).get("chart_palette", [])
    return [hex_to_rgb(c) for c in palette]


def get_font_name(config: dict) -> str:
    return config.get("fonts", {}).get("title", "Calibri")


def format_title(shape, config: dict):
    """Apply standard title formatting to a shape."""
    colors = get_colors(config)
    font_name = get_font_name(config)
    title_size = config.get("fonts", {}).get("title_size_pt", 28)
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.size = Pt(title_size)
        paragraph.font.bold = True
        paragraph.font.color.rgb = colors["primary"]
        paragraph.font.name = font_name


def format_subtitle(shape, config: dict):
    """Apply subtitle formatting."""
    colors = get_colors(config)
    font_name = get_font_name(config)
    subtitle_size = config.get("fonts", {}).get("subtitle_size_pt", 18)
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.size = Pt(subtitle_size)
        paragraph.font.bold = True
        paragraph.font.color.rgb = colors["secondary"]
        paragraph.font.name = font_name


def format_body_text(text_frame, config: dict, size_pt=None):
    """Apply body text formatting to a text frame."""
    colors = get_colors(config)
    font_name = get_font_name(config)
    body_size = size_pt or config.get("fonts", {}).get("body_size_pt", 11)
    for paragraph in text_frame.paragraphs:
        paragraph.font.size = Pt(body_size)
        paragraph.font.color.rgb = colors["text_dark"]
        paragraph.font.name = font_name


def add_slide_title(slide, prs, title_text: str, config: dict):
    """Add a styled title text box to a slide."""
    from pptx.util import Inches, Pt
    txBox = slide.shapes.add_textbox(TITLE_LEFT, TITLE_TOP, TITLE_WIDTH, TITLE_HEIGHT)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(config.get("fonts", {}).get("title_size_pt", 28))
    p.font.bold = True
    p.font.color.rgb = get_colors(config)["primary"]
    p.font.name = get_font_name(config)
    return txBox


def add_accent_bar(slide, config: dict, top=None):
    """Add a colored accent bar below the title."""
    bar_top = top or Inches(1.15)
    colors = get_colors(config)
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        TITLE_LEFT, bar_top,
        Inches(12.3), Inches(0.04),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = colors["accent"]
    shape.line.fill.background()
    return shape


def add_footer(slide, config: dict):
    """Add classification and page footer."""
    colors = get_colors(config)
    font_name = get_font_name(config)
    classification = config.get("classification", "CONFIDENTIAL")
    quarter = config.get("quarter", "Q1")
    year = config.get("year", 2026)

    txBox = slide.shapes.add_textbox(
        CONTENT_LEFT, FOOTER_TOP, Inches(6), FOOTER_HEIGHT
    )
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = f"{classification} — {config.get('company_name', 'TTEC Holdings')} — {quarter} {year}"
    p.font.size = Pt(8)
    p.font.color.rgb = hex_to_rgb("999999")
    p.font.name = font_name


def format_number(value, fmt="money"):
    """Format numbers for display."""
    if value is None:
        return "N/A"
    if fmt == "money":
        if abs(value) >= 1000:
            return f"${value / 1000:.1f}B"
        return f"${value:.0f}M"
    elif fmt == "percent":
        return f"{value:.1f}%"
    elif fmt == "multiple":
        return f"{value:.1f}x"
    elif fmt == "integer":
        return f"{value:,.0f}"
    return str(value)
